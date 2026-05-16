from __future__ import annotations

import logging
import uuid

from celery import shared_task
from django.core.files.base import ContentFile
from django.utils import timezone

logger = logging.getLogger(__name__)


@shared_task(
    bind=True,
    name="apps.cover_letters.tasks.generate_cl_pdf",
    max_retries=3,
    default_retry_delay=30,
    soft_time_limit=120,
    time_limit=180,
    acks_late=True,
)
def generate_cl_pdf(self: object, cover_letter_id: str, user_id: str) -> dict[str, str]:
    """Generate a PDF file for a cover letter and record the download.

    Args:
        cover_letter_id: UUID of the cover letter.
        user_id: UUID of the requesting user.

    Returns:
        A dict with the download id, file_url, and file_size.
    """
    from apps.cover_letters.models import CoverLetter, CoverLetterDownload

    try:
        cover_letter = CoverLetter.objects.select_related("template", "user").get(
            id=cover_letter_id
        )
    except CoverLetter.DoesNotExist:
        logger.error("Cover letter %s not found for PDF generation.", cover_letter_id)
        return {"error": "Cover letter not found."}

    try:
        # Build PDF content from the cover letter data.
        pdf_bytes: bytes = _render_cover_letter_pdf(cover_letter)
        file_name = f"cover_letters/{cover_letter.user_id}/{cover_letter.id}/{uuid.uuid4().hex}.pdf"

        from django.core.files.storage import default_storage

        saved_path: str = default_storage.save(file_name, ContentFile(pdf_bytes))
        file_url: str = default_storage.url(saved_path)
        file_size: int = len(pdf_bytes)

        download = CoverLetterDownload.objects.create(
            cover_letter=cover_letter,
            user_id=user_id,
            format=CoverLetterDownload.Format.PDF,
            file_url=file_url,
            file_size=file_size,
        )

        CoverLetter.objects.filter(id=cover_letter_id).update(
            downloads_count=cover_letter.downloads_count + 1,
            status=CoverLetter.Status.DOWNLOADED,
        )

        logger.info(
            "PDF generated for cover letter %s (download=%s, size=%d bytes).",
            cover_letter_id,
            download.id,
            file_size,
        )
        return {
            "download_id": str(download.id),
            "file_url": file_url,
            "file_size": str(file_size),
        }

    except Exception as exc:
        logger.exception("PDF generation failed for cover letter %s.", cover_letter_id)
        raise self.retry(exc=exc)  # type: ignore[attr-defined]


@shared_task(
    bind=True,
    name="apps.cover_letters.tasks.generate_cl_docx",
    max_retries=3,
    default_retry_delay=30,
    soft_time_limit=120,
    time_limit=180,
    acks_late=True,
)
def generate_cl_docx(self: object, cover_letter_id: str, user_id: str) -> dict[str, str]:
    """Generate a DOCX file for a cover letter and record the download.

    Args:
        cover_letter_id: UUID of the cover letter.
        user_id: UUID of the requesting user.

    Returns:
        A dict with the download id, file_url, and file_size.
    """
    from apps.cover_letters.models import CoverLetter, CoverLetterDownload

    try:
        cover_letter = CoverLetter.objects.select_related("template", "user").get(
            id=cover_letter_id
        )
    except CoverLetter.DoesNotExist:
        logger.error("Cover letter %s not found for DOCX generation.", cover_letter_id)
        return {"error": "Cover letter not found."}

    try:
        docx_bytes: bytes = _render_cover_letter_docx(cover_letter)
        file_name = f"cover_letters/{cover_letter.user_id}/{cover_letter.id}/{uuid.uuid4().hex}.docx"

        from django.core.files.storage import default_storage

        saved_path: str = default_storage.save(file_name, ContentFile(docx_bytes))
        file_url: str = default_storage.url(saved_path)
        file_size: int = len(docx_bytes)

        download = CoverLetterDownload.objects.create(
            cover_letter=cover_letter,
            user_id=user_id,
            format=CoverLetterDownload.Format.DOCX,
            file_url=file_url,
            file_size=file_size,
        )

        CoverLetter.objects.filter(id=cover_letter_id).update(
            downloads_count=cover_letter.downloads_count + 1,
            status=CoverLetter.Status.DOWNLOADED,
        )

        logger.info(
            "DOCX generated for cover letter %s (download=%s, size=%d bytes).",
            cover_letter_id,
            download.id,
            file_size,
        )
        return {
            "download_id": str(download.id),
            "file_url": file_url,
            "file_size": str(file_size),
        }

    except Exception as exc:
        logger.exception("DOCX generation failed for cover letter %s.", cover_letter_id)
        raise self.retry(exc=exc)  # type: ignore[attr-defined]


# ── Private rendering helpers ──────────────────────────────────────


def _render_cover_letter_pdf(cover_letter: object) -> bytes:
    """Render a CoverLetter instance to PDF bytes.

    Uses WeasyPrint to convert an HTML representation of the cover letter
    into a PDF document. Falls back to a simple text-based PDF when
    WeasyPrint is unavailable.
    """
    from apps.cover_letters.models import CoverLetter

    assert isinstance(cover_letter, CoverLetter)
    content: dict = cover_letter.content or {}

    html_parts: list[str] = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        "<style>body{font-family:serif;margin:40px;line-height:1.6;}</style>",
        "</head><body>",
    ]
    if content.get("greeting"):
        html_parts.append(f"<p><strong>{content['greeting']}</strong></p>")
    if content.get("opening"):
        html_parts.append(f"<p>{content['opening']}</p>")
    if content.get("body"):
        html_parts.append(f"<p>{content['body']}</p>")
    if content.get("closing"):
        html_parts.append(f"<p>{content['closing']}</p>")
    if content.get("signature"):
        html_parts.append(f"<p><em>{content['signature']}</em></p>")
    html_parts.append("</body></html>")

    html_string = "".join(html_parts)

    try:
        from weasyprint import HTML  # type: ignore[import-untyped]

        return HTML(string=html_string).write_pdf()
    except ImportError:
        logger.warning("WeasyPrint not installed; generating minimal PDF fallback.")
        return _minimal_text_pdf(_content_to_text(content))


def _render_cover_letter_docx(cover_letter: object) -> bytes:
    """Render a CoverLetter instance to DOCX bytes using python-docx."""
    import io

    from docx import Document  # type: ignore[import-untyped]

    from apps.cover_letters.models import CoverLetter

    assert isinstance(cover_letter, CoverLetter)
    content: dict = cover_letter.content or {}

    doc = Document()
    doc.add_heading(cover_letter.title, level=1)

    for key in ("greeting", "opening", "body", "closing", "signature"):
        text = content.get(key)
        if text:
            doc.add_paragraph(text)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _content_to_text(content: dict) -> str:
    """Flatten cover letter content dict into plain text."""
    parts: list[str] = []
    for key in ("greeting", "opening", "body", "closing", "signature"):
        value = content.get(key)
        if value:
            parts.append(str(value))
    return "\n\n".join(parts)


def _minimal_text_pdf(text: str) -> bytes:
    """Generate a bare-minimum valid PDF containing the given text."""
    lines = text.replace("\r\n", "\n").split("\n")
    stream_lines = ["BT", "/F1 12 Tf"]
    y = 750
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_lines.append(f"1 0 0 1 50 {y} Tm")
        stream_lines.append(f"({safe}) Tj")
        y -= 18
        if y < 50:
            break
    stream_lines.append("ET")
    stream_content = "\n".join(stream_lines)

    objects: list[str] = []
    offsets: list[int] = []
    current = 0

    def add_obj(obj_str: str) -> None:
        nonlocal current
        offsets.append(current)
        objects.append(obj_str)
        current += len(obj_str)

    add_obj("1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n")
    add_obj("2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n")
    add_obj(
        "3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
    )
    add_obj(
        f"4 0 obj\n<< /Length {len(stream_content)} >>\nstream\n{stream_content}\nendstream\nendobj\n"
    )
    add_obj(
        "5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    )

    body = "%PDF-1.4\n" + "".join(objects)
    xref_offset = len(body)
    xref = f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    base = len("%PDF-1.4\n")
    for off in offsets:
        xref += f"{base + off:010d} 00000 n \n"
    xref += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"

    return (body + xref).encode("latin-1")
