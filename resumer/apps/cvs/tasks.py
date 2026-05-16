from __future__ import annotations

import logging
import uuid

from celery import shared_task
from django.core.files.base import ContentFile
from django.utils import timezone

logger = logging.getLogger(__name__)


@shared_task(
    bind=True,
    name="apps.cvs.tasks.generate_cv_pdf",
    max_retries=3,
    default_retry_delay=30,
    soft_time_limit=120,
    time_limit=180,
    acks_late=True,
)
def generate_cv_pdf(self: object, cv_id: str, user_id: str) -> dict[str, str]:
    """Generate a PDF file for a CV and record the download.

    Args:
        cv_id: UUID of the CV.
        user_id: UUID of the requesting user.

    Returns:
        A dict with the download id, file_url, and file_size.
    """
    from apps.cvs.models import CV, CVDownload

    try:
        cv = CV.objects.select_related("template", "user").prefetch_related("sections").get(id=cv_id)
    except CV.DoesNotExist:
        logger.error("CV %s not found for PDF generation.", cv_id)
        return {"error": "CV not found."}

    try:
        pdf_bytes: bytes = _render_cv_pdf(cv)
        file_name = f"cvs/{cv.user_id}/{cv.id}/{uuid.uuid4().hex}.pdf"

        from django.core.files.storage import default_storage

        saved_path: str = default_storage.save(file_name, ContentFile(pdf_bytes))
        file_url: str = default_storage.url(saved_path)
        file_size: int = len(pdf_bytes)

        download = CVDownload.objects.create(
            cv=cv,
            user_id=user_id,
            format=CVDownload.Format.PDF,
            file_url=file_url,
            file_size=file_size,
        )

        CV.objects.filter(id=cv_id).update(
            downloads_count=cv.downloads_count + 1,
            status=CV.Status.DOWNLOADED,
        )

        logger.info(
            "PDF generated for CV %s (download=%s, size=%d bytes).",
            cv_id,
            download.id,
            file_size,
        )
        return {
            "download_id": str(download.id),
            "file_url": file_url,
            "file_size": str(file_size),
        }

    except Exception as exc:
        logger.exception("PDF generation failed for CV %s.", cv_id)
        raise self.retry(exc=exc)  # type: ignore[attr-defined]


@shared_task(
    bind=True,
    name="apps.cvs.tasks.generate_cv_docx",
    max_retries=3,
    default_retry_delay=30,
    soft_time_limit=120,
    time_limit=180,
    acks_late=True,
)
def generate_cv_docx(self: object, cv_id: str, user_id: str) -> dict[str, str]:
    """Generate a DOCX file for a CV and record the download.

    Args:
        cv_id: UUID of the CV.
        user_id: UUID of the requesting user.

    Returns:
        A dict with the download id, file_url, and file_size.
    """
    from apps.cvs.models import CV, CVDownload

    try:
        cv = CV.objects.select_related("template", "user").prefetch_related("sections").get(id=cv_id)
    except CV.DoesNotExist:
        logger.error("CV %s not found for DOCX generation.", cv_id)
        return {"error": "CV not found."}

    try:
        docx_bytes: bytes = _render_cv_docx(cv)
        file_name = f"cvs/{cv.user_id}/{cv.id}/{uuid.uuid4().hex}.docx"

        from django.core.files.storage import default_storage

        saved_path: str = default_storage.save(file_name, ContentFile(docx_bytes))
        file_url: str = default_storage.url(saved_path)
        file_size: int = len(docx_bytes)

        download = CVDownload.objects.create(
            cv=cv,
            user_id=user_id,
            format=CVDownload.Format.DOCX,
            file_url=file_url,
            file_size=file_size,
        )

        CV.objects.filter(id=cv_id).update(
            downloads_count=cv.downloads_count + 1,
            status=CV.Status.DOWNLOADED,
        )

        logger.info(
            "DOCX generated for CV %s (download=%s, size=%d bytes).",
            cv_id,
            download.id,
            file_size,
        )
        return {
            "download_id": str(download.id),
            "file_url": file_url,
            "file_size": str(file_size),
        }

    except Exception as exc:
        logger.exception("DOCX generation failed for CV %s.", cv_id)
        raise self.retry(exc=exc)  # type: ignore[attr-defined]


# ── Private rendering helpers ──────────────────────────────────────


def _render_cv_pdf(cv: object) -> bytes:
    """Render a CV instance (with its sections) to PDF bytes.

    Uses WeasyPrint when available; falls back to a minimal text PDF.
    """
    from apps.cvs.models import CV

    assert isinstance(cv, CV)
    content: dict = cv.content or {}
    sections = cv.sections.filter(is_visible=True).order_by("order")

    html_parts: list[str] = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        "<style>body{font-family:serif;margin:40px;line-height:1.6;} "
        "h1{margin-bottom:4px;} h2{margin-top:20px;border-bottom:1px solid #ccc;padding-bottom:4px;}</style>",
        "</head><body>",
        f"<h1>{cv.title}</h1>",
    ]

    # Render top-level content JSON fields.
    if content.get("summary"):
        html_parts.append(f"<p>{content['summary']}</p>")

    # Render each visible section.
    for section in sections:
        section_title = section.get_section_type_display()
        html_parts.append(f"<h2>{section_title}</h2>")
        section_content = section.content or {}
        if isinstance(section_content, dict):
            for key, value in section_content.items():
                if isinstance(value, list):
                    html_parts.append("<ul>")
                    for item in value:
                        html_parts.append(f"<li>{item}</li>")
                    html_parts.append("</ul>")
                else:
                    html_parts.append(f"<p><strong>{key}:</strong> {value}</p>")
        elif isinstance(section_content, str):
            html_parts.append(f"<p>{section_content}</p>")

    html_parts.append("</body></html>")
    html_string = "".join(html_parts)

    try:
        from weasyprint import HTML  # type: ignore[import-untyped]

        return HTML(string=html_string).write_pdf()
    except ImportError:
        logger.warning("WeasyPrint not installed; generating minimal PDF fallback.")
        return _minimal_text_pdf(_cv_to_text(cv, sections))


def _render_cv_docx(cv: object) -> bytes:
    """Render a CV instance to DOCX bytes using python-docx."""
    import io

    from docx import Document  # type: ignore[import-untyped]

    from apps.cvs.models import CV

    assert isinstance(cv, CV)
    content: dict = cv.content or {}
    sections = cv.sections.filter(is_visible=True).order_by("order")

    doc = Document()
    doc.add_heading(cv.title, level=1)

    if content.get("summary"):
        doc.add_paragraph(content["summary"])

    for section in sections:
        doc.add_heading(section.get_section_type_display(), level=2)
        section_content = section.content or {}
        if isinstance(section_content, dict):
            for key, value in section_content.items():
                if isinstance(value, list):
                    for item in value:
                        doc.add_paragraph(str(item), style="List Bullet")
                else:
                    doc.add_paragraph(f"{key}: {value}")
        elif isinstance(section_content, str):
            doc.add_paragraph(section_content)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _cv_to_text(cv: object, sections: object) -> str:
    """Flatten CV data into plain text for the fallback PDF renderer."""
    from apps.cvs.models import CV

    assert isinstance(cv, CV)
    content: dict = cv.content or {}
    parts: list[str] = [cv.title, ""]

    if content.get("summary"):
        parts.append(content["summary"])
        parts.append("")

    for section in sections:  # type: ignore[union-attr]
        parts.append(f"--- {section.get_section_type_display()} ---")
        sc = section.content or {}
        if isinstance(sc, dict):
            for key, value in sc.items():
                if isinstance(value, list):
                    for item in value:
                        parts.append(f"  - {item}")
                else:
                    parts.append(f"  {key}: {value}")
        elif isinstance(sc, str):
            parts.append(f"  {sc}")
        parts.append("")

    return "\n".join(parts)


def _minimal_text_pdf(text: str) -> bytes:
    """Generate a bare-minimum valid PDF containing the given text."""
    lines = text.replace("\r\n", "\n").split("\n")
    stream_lines = ["BT", "/F1 12 Tf"]
    y = 750
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_lines.append(f"1 0 0 1 50 {y} Tm")
        stream_lines.append(f"({safe}) Tj")
        y -= 18
        if y < 50:
            break
    stream_lines.append("ET")
    stream_content = "\n".join(stream_lines)

    objects: list[str] = []
    offsets: list[int] = []
    current = 0

    def add_obj(obj_str: str) -> None:
        nonlocal current
        offsets.append(current)
        objects.append(obj_str)
        current += len(obj_str)

    add_obj("1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n")
    add_obj("2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n")
    add_obj(
        "3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
    )
    add_obj(
        f"4 0 obj\n<< /Length {len(stream_content)} >>\nstream\n{stream_content}\nendstream\nendobj\n"
    )
    add_obj(
        "5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    )

    body = "%PDF-1.4\n" + "".join(objects)
    xref_offset = len(body)
    xref = f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    base = len("%PDF-1.4\n")
    for off in offsets:
        xref += f"{base + off:010d} 00000 n \n"
    xref += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"

    return (body + xref).encode("latin-1")
